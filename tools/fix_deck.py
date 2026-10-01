"""Post-conversion fix-ups for the C716 deck: strip the WSQ layer convert_deck.py left behind."""
import copy, re
from pptx import Presentation

PATH = "courseware/Semiconductor Process Fundamentals Masterclass (C716)-v1.0.pptx"
SRC = ("source-wsq/courseware/WSQ - Master Trainer Slides - Mastering Semiconductor Key Modules for Enhanced "
       "CMOS Process Integration and Device Performance - v9.0.pptx")
p = Presentation(PATH)
src = Presentation(SRC)

def shape(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise SystemExit(f"shape {sid} not found")

def set_para(para, line):
    runs = [r for r in para.runs if r.text] or para.runs
    runs[0].text = line
    for r in para.runs:
        if r._r is not runs[0]._r:
            r.text = ""

def set_paras(slide, sid, lines):
    paras = shape(slide, sid).text_frame.paragraphs
    assert len(paras) == len(lines), (sid, len(paras), lines)
    for para, line in zip(paras, lines):
        set_para(para, line)

def find(title):
    hits = [s for s in p.slides if any(sh.has_text_frame and sh.text_frame.text == title for sh in s.shapes)]
    assert len(hits) == 1, (title, len(hits))
    return hits[0]

# ---- Course Profile
prof = find("Course Profile")
set_paras(prof, 11, ["2 days · 15 hours"])
set_paras(prof, 14, ["Course code: C716"])
set_paras(prof, 17, ["9:30am–5:30pm · 7.5 hours per day"])
set_paras(prof, 20, ["Topic 1: process flow"])
set_paras(prof, 22, ["Evidence and practice"])
set_paras(prof, 26, ["Topic 2: integration and performance"])
set_paras(prof, 32, ["Worksheets, checklists and evidence records"])
set_paras(prof, 35, ["Course synthesis on Day 2"])

# ---- Lesson Plan slides: 9:30am-5:30pm, one 30-min lunch (blocks match the LP schedule tables)
d1 = find("Lesson Plan - Day 1"); d2 = find("Lesson Plan - Day 2")
for sid, line in [(11, "09:30 Introductions and course overview"), (14, "09:50 Block 1A: patterning and etch"),
                  (17, "10:55 Activities 1-2"), (20, "12:40 Lunch"),
                  (26, "13:10 Block 1B: implant, thermal and films"), (29, "14:30 Activities 3-4"),
                  (32, "16:20 Integrated-flow review"), (35, "17:30 End of Day 1")]:
    set_paras(d1, sid, [line])
for sid, line in [(11, "09:30 Day 1 retrieval; Block 2A: performance, yield and reliability"), (14, "11:20 Activity 5"),
                  (17, "12:20 Lunch"), (20, "12:50 Activity 6"),
                  (22, "Afternoon"), (26, "14:05 Block 2B: SPC and packaging"),
                  (29, "15:20 Activities 7-8"), (32, "17:00 Course synthesis and feedback"),
                  (35, "17:30 End of class")]:
    set_paras(d2, sid, [line])

# ---- Activity meta lines: "LO1 · K1 / A1 / A2 · 45 minutes" -> "LO1 · <LP minutes>"
MINUTES = {1: 45, 2: 60, 3: 55, 4: 55, 5: 60, 6: 75, 7: 50, 8: 50}
META = re.compile(r"^(LO\d)(?:\s*·\s*[KA]\d+(?:\s*/\s*[KA]\d+)*)\s*·\s*\d+ minutes$")
act = 0
for s in p.slides:
    texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
    if any(re.fullmatch(r"ACTIVITY \d", t) for t in texts):
        act = int(next(t for t in texts if re.fullmatch(r"ACTIVITY \d", t)).split()[1])
    for sh in s.shapes:
        if sh.has_text_frame:
            mo = META.match(sh.text_frame.text)
            if mo:
                set_paras(s, sh.shape_id, [f"{mo.group(1)} · {MINUTES[act]} minutes"])

# ---- activity goals that referenced the Practical Performance tasks
PP = {"Produce the process-flow and challenge analysis used as the foundation for Practical Performance Task 1.":
          "Produce the process-flow and challenge analysis for a shallow trench isolation integration.",
      "Recreate the legacy P-chart/SPC exercises and generate the corrective-action chain used in Practical Performance Task 2.":
          "Recreate the legacy P-chart/SPC exercises and generate the corrective-action chain for a CD excursion."}
for s in p.slides:
    for sh in s.shapes:
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                t = "".join(r.text for r in para.runs)
                if t in PP:
                    set_para(para, PP[t])

# ---- "OFFICIAL TOPIC" (WSQ syllabus framing) -> "TOPIC"
for s in p.slides:
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                if "OFFICIAL TOPIC" in r.text:
                    r.text = r.text.replace("OFFICIAL TOPIC", "TOPIC")

# ---- Course Feedback: restore the digital-attendance slide's text from the source, then rewrite in place
fb = p.slides[-2]
srcfb = src.slides[137]
for sid in (3, 4, 9, 13, 17, 21):
    dst, s0 = shape(fb, sid), shape(srcfb, sid)
    dst._element.replace(dst._element.txBody, copy.deepcopy(s0._element.txBody))
set_paras(fb, 3, ["COURSE FEEDBACK"])
set_paras(fb, 4, ["Course Feedback & Certificate"])
set_paras(fb, 9, ["Complete the course feedback survey before you leave"])
set_paras(fb, 13, ["Tell us what worked and what to improve"])
set_paras(fb, 17, ["Ask the trainer if you cannot open the survey"])
set_paras(fb, 21, ["Your certificate of completion will be emailed to you"])

# ---- hyperlink relationships still pointing at the WSQ repo
for s in p.slides:
    for rel in s.part.rels.values():
        if rel.is_external and "TGS-2024051248" in rel._target:
            rel._target = rel._target.replace(
                "TGS-2024051248-Mastering-Semiconductor-Key-Modules-for-Enhanced-CMOS-Process-Integration-and-Device",
                "C716-Semiconductor-Process-Fundamentals-Masterclass")
p.core_properties.title = "Semiconductor Process Fundamentals Masterclass"
p.core_properties.subject = "C716"
p.save(PATH)
print("saved")
